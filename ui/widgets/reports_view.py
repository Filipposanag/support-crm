# -*- coding: utf-8 -*-
"""
Reports View - Οθόνη Αναφορών (Support CRM)
A clean, self-contained financial & activity report per period.
Reads live data via db_connection.execute_query() (same pattern as
DashboardView / DebtsView). Replaces the legacy family/subscription report.
"""

import logging
from datetime import date

from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QGridLayout, QLabel, QFrame,
    QPushButton, QDateEdit, QComboBox, QTableWidget, QTableWidgetItem,
    QHeaderView, QAbstractItemView, QMessageBox, QFileDialog,
)
from PySide6.QtCore import Qt, QDate
from PySide6.QtGui import QColor, QBrush
from ui.utils.permission_helper import has_permission

logger = logging.getLogger(__name__)

_PLAN = {"hourly": "Ωριαία", "monthly": "Μηνιαία",
         "yearly": "Ετήσια", "custom": "Προσαρμοσμένο"}


def _fmt_hours(mins):
    try:
        tm = int(round(float(mins or 0)))
    except (TypeError, ValueError):
        tm = 0
    hh, mm = divmod(tm, 60)
    if hh and mm:
        return f"{hh}ω {mm}λ"
    if hh:
        return f"{hh}ω"
    return f"{mm}λ"


class ReportsView(QWidget):
    """Period financial & activity report."""

    def __init__(self, db_connection, tr_manager=None, parent=None):
        super().__init__(parent)
        self.db = db_connection
        self.tr_manager = tr_manager
        self._kpi = {}
        self._rows = []
        self._build_ui()
        self.refresh()
        logger.debug("ReportsView initialized")

    # ---------- data ----------
    def _query(self, sql, params=None, default=None):
        try:
            rows = self.db.execute_query(sql, params)
            return rows if rows is not None else (default or [])
        except Exception as e:
            logger.error(f"Reports query failed: {e}")
            return default if default is not None else []

    # ---------- ui ----------
    def _build_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(24, 18, 24, 18)
        root.setSpacing(14)

        # Header
        header = QHBoxLayout()
        title = QLabel("Αναφορές")
        title.setStyleSheet("font-size:22px;font-weight:bold;color:#1976D2;")
        header.addWidget(title)
        header.addStretch()
        self.export_btn = QPushButton("Εξαγωγή σε Excel")
        self.export_btn.setStyleSheet("QPushButton{background-color:#4CAF50;color:white;font-weight:bold;padding:8px 16px;border:none;border-radius:4px;}QPushButton:hover{background-color:#45a049;}QPushButton:disabled{background-color:#cccccc;color:#8a8a8a;}")
        self.export_btn.clicked.connect(self._on_export)
        self.export_btn.setEnabled(has_permission('reports.export'))
        header.addWidget(self.export_btn)
        root.addLayout(header)

        # Period filter bar
        fbar = QHBoxLayout()
        fbar.setSpacing(8)
        fbar.addWidget(QLabel("Περίοδος:"))
        self.preset_combo = QComboBox()
        self.preset_combo.addItem("Τρέχον έτος", "year")
        self.preset_combo.addItem("Τελευταίοι 12 μήνες", "12m")
        self.preset_combo.addItem("Τρέχων μήνας", "month")
        self.preset_combo.addItem("Προσαρμοσμένη", "custom")
        self.preset_combo.currentIndexChanged.connect(self._on_preset)
        fbar.addWidget(self.preset_combo)

        fbar.addWidget(QLabel("Από:"))
        self.from_edit = QDateEdit()
        self.from_edit.setCalendarPopup(True)
        self.from_edit.setDisplayFormat("dd/MM/yyyy")
        self.from_edit.setDate(QDate(QDate.currentDate().year(), 1, 1))
        fbar.addWidget(self.from_edit)

        fbar.addWidget(QLabel("Έως:"))
        self.to_edit = QDateEdit()
        self.to_edit.setCalendarPopup(True)
        self.to_edit.setDisplayFormat("dd/MM/yyyy")
        self.to_edit.setDate(QDate.currentDate())
        fbar.addWidget(self.to_edit)

        self.apply_btn = QPushButton("Ανανέωση")
        self.apply_btn.setStyleSheet("QPushButton{background-color:#1976D2;color:white;font-weight:bold;padding:8px 16px;border:none;border-radius:4px;}QPushButton:hover{background-color:#1565C0;}QPushButton:disabled{background-color:#cccccc;color:#8a8a8a;}")
        self.apply_btn.clicked.connect(self.refresh)
        fbar.addWidget(self.apply_btn)
        fbar.addStretch()
        root.addLayout(fbar)

        # KPI cards
        kpi = QGridLayout()
        kpi.setSpacing(12)
        cards = [
            ("clients", "Πελάτες με Κίνηση", "#1976D2"),
            ("sessions", "Συνεδρίες", "#00838F"),
            ("hours", "Ώρες Υποστήριξης", "#6A1B9A"),
            ("charged", "Σύνολο Χρεώσεων", "#455A64"),
            ("paid", "Εισπράξεις", "#2E7D32"),
            ("debt", "Οφειλές", "#C62828"),
        ]
        for i, (k, lbl, col) in enumerate(cards):
            kpi.addWidget(self._make_card(k, lbl, col), i // 3, i % 3)
        root.addLayout(kpi)

        # Table
        self.table = QTableWidget()
        self.table.setColumnCount(7)
        self.table.setHorizontalHeaderLabels([
            "Κωδικός", "Πελάτης", "Συνεδρίες", "Ώρες",
            "Χρεώσεις €", "Εισπράξεις €", "Οφειλή €"])
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SingleSelection)
        self.table.setAlternatingRowColors(True)
        self.table.verticalHeader().setVisible(False)
        self.table.setSortingEnabled(True)
        hh = self.table.horizontalHeader()
        hh.setSectionResizeMode(1, QHeaderView.Stretch)
        for c in (0, 2, 3, 4, 5, 6):
            hh.setSectionResizeMode(c, QHeaderView.ResizeToContents)
        root.addWidget(self.table, 1)

        # Footer totals
        self.footer_lbl = QLabel("")
        self.footer_lbl.setStyleSheet("font-size:13px;color:#555;")
        root.addWidget(self.footer_lbl)

    def _make_card(self, key, label, color):
        card = QFrame()
        card.setStyleSheet(
            f"QFrame{{background:white;border:1px solid #E0E0E0;"
            f"border-left:5px solid {color};border-radius:8px;}}")
        card.setMinimumHeight(80)
        v = QVBoxLayout(card)
        v.setContentsMargins(16, 10, 16, 10)
        val = QLabel("—")
        val.setStyleSheet(f"font-size:22px;font-weight:bold;color:{color};border:none;")
        lb = QLabel(label)
        lb.setStyleSheet("font-size:12px;color:#555;border:none;")
        lb.setWordWrap(True)
        v.addWidget(val)
        v.addWidget(lb)
        self._kpi[key] = val
        return card

    # ---------- period ----------
    def _on_preset(self):
        mode = self.preset_combo.currentData()
        today = QDate.currentDate()
        custom = (mode == "custom")
        self.from_edit.setEnabled(custom)
        self.to_edit.setEnabled(custom)
        if mode == "year":
            self.from_edit.setDate(QDate(today.year(), 1, 1))
            self.to_edit.setDate(today)
        elif mode == "12m":
            self.from_edit.setDate(today.addMonths(-12))
            self.to_edit.setDate(today)
        elif mode == "month":
            self.from_edit.setDate(QDate(today.year(), today.month(), 1))
            self.to_edit.setDate(today)
        if mode != "custom":
            self.refresh()

    def _period(self):
        d_from = self.from_edit.date().toString("yyyy-MM-dd")
        d_to = self.to_edit.date().toString("yyyy-MM-dd")
        return d_from, d_to

    # ---------- load ----------
    def refresh(self):
        d_from, d_to = self._period()

        sess = self._query(
            "SELECT client_id cid, COUNT(*) n, "
            "COALESCE(SUM(CASE WHEN status='completed' THEN duration_minutes ELSE 0 END),0) mins "
            "FROM support_sessions "
            "WHERE DATE(start_time) BETWEEN %s AND %s "
            "GROUP BY client_id", (d_from, d_to), default=[])
        pays = self._query(
            "SELECT family_id cid, COALESCE(SUM(total_amount),0) tot, "
            "COALESCE(SUM(paid_amount),0) paid, COALESCE(SUM(remaining_amount),0) rem "
            "FROM payments WHERE payment_date BETWEEN %s AND %s "
            "GROUP BY family_id", (d_from, d_to), default=[])
        clients = self._query(
            "SELECT id, legacy_code code, last_name, first_name, company_name "
            "FROM families", default=[])

        cmap = {c["id"]: c for c in clients}
        smap = {r["cid"]: r for r in sess}
        pmap = {r["cid"]: r for r in pays}

        ids = set(smap) | set(pmap)
        rows = []
        for cid in ids:
            c = cmap.get(cid, {})
            s = smap.get(cid, {})
            p = pmap.get(cid, {})
            name = f"{c.get('last_name','') or ''} {c.get('first_name','') or ''}".strip()
            comp = (c.get("company_name") or "").strip()
            if comp:
                name = f"{name} ({comp})" if name else comp
            rows.append({
                "cid": cid,
                "code": c.get("code", ""),
                "name": name or "—",
                "sessions": int(s.get("n", 0) or 0),
                "mins": float(s.get("mins", 0) or 0),
                "charged": float(p.get("tot", 0) or 0),
                "paid": float(p.get("paid", 0) or 0),
                "debt": float(p.get("rem", 0) or 0),
            })
        rows.sort(key=lambda x: (x["debt"], x["charged"]), reverse=True)
        self._rows = rows
        self._populate(rows)
        self._update_kpis(rows)

    def _populate(self, rows):
        self.table.setSortingEnabled(False)
        self.table.setRowCount(0)
        for r in rows:
            i = self.table.rowCount()
            self.table.insertRow(i)
            cells = [
                (str(r["code"]), Qt.AlignCenter, None),
                (r["name"], Qt.AlignVCenter | Qt.AlignLeft, None),
                (str(r["sessions"]), Qt.AlignCenter, None),
                (_fmt_hours(r["mins"]), Qt.AlignCenter, None),
                (f"{r['charged']:.2f}", Qt.AlignRight | Qt.AlignVCenter, None),
                (f"{r['paid']:.2f}", Qt.AlignRight | Qt.AlignVCenter, "#2E7D32"),
                (f"{r['debt']:.2f}", Qt.AlignRight | Qt.AlignVCenter,
                 "#C62828" if r["debt"] > 0 else None),
            ]
            for col, (text, align, color) in enumerate(cells):
                it = QTableWidgetItem()
                # numeric sort for numeric columns
                if col in (2, 4, 5, 6):
                    it.setData(Qt.EditRole, float(
                        r["sessions"] if col == 2 else
                        (r["charged"] if col == 4 else
                         (r["paid"] if col == 5 else r["debt"]))))
                    it.setText(text)
                elif col == 3:
                    it.setData(Qt.EditRole, r["mins"])
                    it.setText(text)
                else:
                    it.setText(text)
                it.setTextAlignment(align)
                if col == 0:
                    it.setData(Qt.UserRole, r["cid"])
                if color:
                    f = it.font(); f.setBold(True); it.setFont(f)
                    it.setForeground(QBrush(QColor(color)))
                self.table.setItem(i, col, it)
        self.table.setSortingEnabled(True)

    def _update_kpis(self, rows):
        n_clients = len(rows)
        sessions = sum(r["sessions"] for r in rows)
        mins = sum(r["mins"] for r in rows)
        charged = sum(r["charged"] for r in rows)
        paid = sum(r["paid"] for r in rows)
        debt = sum(r["debt"] for r in rows)
        hh, mm = divmod(int(round(mins)), 60)
        htxt = (f"{hh}ω {mm}λ" if hh and mm else (f"{hh}ω" if hh else f"{mm}λ"))
        self._kpi["clients"].setText(str(n_clients))
        self._kpi["sessions"].setText(str(sessions))
        self._kpi["hours"].setText(htxt)
        self._kpi["charged"].setText(f"{charged:.2f} €")
        self._kpi["paid"].setText(f"{paid:.2f} €")
        self._kpi["debt"].setText(f"{debt:.2f} €")
        d_from, d_to = self._period()
        self.footer_lbl.setText(
            f"Περίοδος {self.from_edit.date().toString('dd/MM/yyyy')} – "
            f"{self.to_edit.date().toString('dd/MM/yyyy')}  ·  "
            f"{n_clients} πελάτες  ·  {sessions} συνεδρίες  ·  "
            f"Χρεώσεις {charged:.2f}€  ·  Εισπράξεις {paid:.2f}€  ·  Οφειλές {debt:.2f}€")

    # ---------- export ----------
    def _on_export(self):
        if not has_permission('reports.export'):
            QMessageBox.warning(self, "Πρόσβαση Απορρίφθηκε",
                "Δεν έχετε δικαίωμα εξαγωγής αναφορών.")
            return
        if not self._rows:
            QMessageBox.information(self, "Εξαγωγή",
                                    "Δεν υπάρχουν δεδομένα για εξαγωγή.")
            return
        default = f"anafora_{date.today().strftime('%Y%m%d')}.xlsx"
        path, _ = QFileDialog.getSaveFileName(
            self, "Αποθήκευση Excel", default, "Excel (*.xlsx)")
        if not path:
            return
        try:
            self._export_xlsx(path)
            QMessageBox.information(self, "Εξαγωγή",
                                   f"Η εξαγωγή ολοκληρώθηκε:\n{path}")
        except Exception as e:
            logger.error(f"Report export failed: {e}", exc_info=True)
            QMessageBox.warning(self, "Σφάλμα", f"Αποτυχία εξαγωγής:\n{e}")

    def _export_xlsx(self, path):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = Workbook()
        ws = wb.active
        ws.title = "Αναφορά"
        d_from = self.from_edit.date().toString("dd/MM/yyyy")
        d_to = self.to_edit.date().toString("dd/MM/yyyy")
        ws.append([f"Αναφορά περιόδου {d_from} – {d_to}"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([])
        headers = ["Κωδικός", "Πελάτης", "Συνεδρίες", "Ώρες (λεπτά)",
                   "Χρεώσεις €", "Εισπράξεις €", "Οφειλή €"]
        ws.append(headers)
        hrow = ws.max_row
        fill = PatternFill("solid", fgColor="1976D2")
        for cell in ws[hrow]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")

        tot_s = tot_c = tot_p = tot_d = 0
        for r in self._rows:
            ws.append([r["code"], r["name"], r["sessions"], round(r["mins"], 0),
                       round(r["charged"], 2), round(r["paid"], 2), round(r["debt"], 2)])
            tot_s += r["sessions"]; tot_c += r["charged"]
            tot_p += r["paid"]; tot_d += r["debt"]
        ws.append([])
        ws.append(["", "ΣΥΝΟΛΑ", tot_s, "", round(tot_c, 2),
                   round(tot_p, 2), round(tot_d, 2)])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)

        widths = [10, 34, 12, 14, 14, 14, 14]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + i)].width = w
        wb.save(path)

    # ---------- misc (compat) ----------
    def update_translations(self):
        pass

    def apply_column_preferences(self):
        pass
